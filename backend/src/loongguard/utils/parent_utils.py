"""
家长端工具函数（JWT 签发校验 / 密码哈希 / Excel 导入导出）

来源：分仓库 ParentModel/website/utils.py，环境变量统一为 LG_PARENT_ 前缀。
密码哈希沿用 passlib pbkdf2_sha256 格式，与既有 class_monitor.db 中
管理员口令哈希兼容，可直接沿用原账号数据。
"""

from __future__ import annotations

import hashlib
import io
import logging
import os
import time
from pathlib import Path

import jwt
from openpyxl import Workbook, load_workbook
from passlib.hash import pbkdf2_sha256

logger = logging.getLogger(__name__)

_BACKEND_DIR = Path(__file__).resolve().parents[3]

# JWT 有效期（秒）。旧版本签发的无 exp token 在过期机制引入后仍可解码
# （PyJWT 仅在 token 携带 exp 时才校验），但新签发 token 一律带过期。
TOKEN_TTL_SEC = 7 * 24 * 3600

# 生产环境禁止使用的占位密钥（与 .env.example 中的示例值一致）
_PLACEHOLDER_SECRETS = ("change-me-admin-jwt-secret", "change-me-wx-jwt-secret")

# 注意：.env 由 pipeline.main() 在运行时加载（晚于模块导入），
# 所有环境变量必须惰性读取，不能在模块级固化。


def _admin_jwt_secret() -> str:
    return os.environ.get("LG_PARENT_ADMIN_JWT_SECRET", "")


def _wx_jwt_secret() -> str:
    return os.environ.get("LG_PARENT_WX_JWT_SECRET", "")


def template_dir() -> str:
    """Excel 导入模板生成目录（管理后台下载用）"""
    return os.environ.get("LG_PARENT_TEMPLATE_DIR") or str(
        _BACKEND_DIR / "static" / "template"
    )


def upload_dir() -> str:
    """家长端上传文件根目录（成长记录图片等）"""
    d = os.environ.get("LG_PARENT_UPLOAD_DIR") or str(
        _BACKEND_DIR / "data" / "parent" / "uploads"
    )
    os.makedirs(d, exist_ok=True)
    return d


def hash_pwd(pwd: str) -> str:
    return pbkdf2_sha256.hash(pwd)


def verify_pwd(raw: str, hash_str: str) -> bool:
    return pbkdf2_sha256.verify(raw, hash_str)


def generate_wx_token(openid: str) -> str:
    now = int(time.time())
    return jwt.encode(
        {"openid": openid, "iat": now, "exp": now + TOKEN_TTL_SEC},
        _wx_jwt_secret(), algorithm="HS256",
    )


def decode_wx_token(token: str) -> dict | None:
    """校验家长 JWT，失败返回 None（由调用方决定 401 响应）

    PyJWT 在 token 携带 exp 时自动校验过期，过期即抛出并返回 None。
    """
    try:
        return jwt.decode(token, _wx_jwt_secret(), algorithms=["HS256"])
    except Exception:
        return None


def generate_admin_token(username: str) -> str:
    now = int(time.time())
    return jwt.encode(
        {"username": username, "iat": now, "exp": now + TOKEN_TTL_SEC},
        _admin_jwt_secret(), algorithm="HS256",
    )


def decode_admin_token(token: str) -> dict | None:
    """校验管理员 JWT，失败返回 None"""
    try:
        return jwt.decode(token, _admin_jwt_secret(), algorithms=["HS256"])
    except Exception:
        return None


def validate_jwt_secrets() -> None:
    """
    启动期校验家长端两把 JWT 密钥。

    空密钥意味着 HS256 用空串签发，任何人可伪造任意家长/管理员 token；
    占位密钥等同于公开密钥。development/test 仅记录 error 日志，
    production 直接阻断启动。
    """
    checks = (
        ("LG_PARENT_ADMIN_JWT_SECRET", _admin_jwt_secret()),
        ("LG_PARENT_WX_JWT_SECRET", _wx_jwt_secret()),
    )
    is_prod = os.environ.get("LG_ENV", "").strip() == "production"
    for name, val in checks:
        problem = ""
        if not val.strip():
            problem = "未设置（空密钥签发的 token 可被任意伪造）"
        elif val.strip() in _PLACEHOLDER_SECRETS:
            problem = "仍为 .env.example 占位值"
        if not problem:
            continue
        msg = f"{name} {problem}，请配置为 `openssl rand -hex 32` 生成的随机密钥"
        if is_prod:
            raise RuntimeError(f"生产环境安全校验失败: {msg}")
        logger.error("%s（LG_ENV=production 时将阻断启动）", msg)


def is_local_network(ip: str) -> bool:
    local_prefix = ("127.", "192.168.", "10.", "172.16.")
    return ip.startswith(local_prefix)


def generate_openid(phone: str) -> str:
    """后台批量导入时为无微信家长生成占位 openid"""
    return "import_" + hashlib.md5(phone.encode()).hexdigest()[:16]


def create_data_import_template() -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "数据导入"
    headers = ["学生姓名", "班级名称", "年级", "家长姓名", "电话", "远程观看权限", "家长OpenID（可选）"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    ws.append(["张三", "中一班", 2026, "张三爸爸", "13800138001", 1, ""])
    ws.freeze_panes = "A2"
    os.makedirs(template_dir(), exist_ok=True)
    path = os.path.join(template_dir(), "数据导入模板.xlsx")
    wb.save(path)
    return path


def _num_str(v) -> str | None:
    """Excel 数值单元格经 openpyxl 常读到 float（2026.0 / 1.38e+10），
    统一转成无科学计数法的整数字符串，避免手机号/年级被误判非法"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse_data_import_excel(file_path: str) -> list[dict]:
    wb = load_workbook(file_path)
    ws = wb.active
    items: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        student_name = row[0] if len(row) > 0 else None
        class_name = row[1] if len(row) > 1 else None
        year = row[2] if len(row) > 2 else None
        if isinstance(year, float) and year.is_integer():
            year = int(year)
        parent_name = row[3] if len(row) > 3 else None
        phone = _num_str(row[4]) if len(row) > 4 and row[4] is not None else None
        watch_switch = row[5] if len(row) > 5 else 1
        openid = str(row[6]) if len(row) > 6 and row[6] is not None else None

        errors: list[str] = []
        if not student_name:
            errors.append("学生姓名缺失")
        if not class_name:
            errors.append("班级名称缺失")
        if not year:
            errors.append("年级缺失")
        elif not str(year).isdigit():
            errors.append("年级必须为数字")
        if not parent_name:
            errors.append("家长姓名缺失")
        if not phone:
            errors.append("家长电话缺失")

        try:
            year_int = int(year) if year else None
        except (TypeError, ValueError):
            year_int = None
            if "年级必须为数字" not in errors:
                errors.append("年级必须为数字")
        try:
            watch_int = int(watch_switch) if watch_switch is not None else 1
        except (TypeError, ValueError):
            watch_int = 1
            errors.append("远程观看权限必须为数字，已默认为1")

        items.append({
            "student_name": str(student_name).strip() if student_name else "",
            "class_name": str(class_name).strip() if class_name else "",
            "year": year_int,
            "parent_name": str(parent_name).strip() if parent_name else "",
            "phone": str(phone).strip() if phone else "",
            "watch_switch": watch_int,
            "openid": str(openid).strip() if openid else None,
            "errors": errors,
        })
    return items


# ==================== 导出函数 ====================
def export_students_with_parents() -> bytes:
    from loongguard.db.parent_db import get_student_parent_bindings
    data = get_student_parent_bindings()
    wb = Workbook()
    ws = wb.active
    ws.title = "学生信息"
    headers = ["学生ID", "学生姓名", "班级", "年级", "远程观看权限",
               "家长OpenID", "家长姓名", "家长手机号", "家长昵称"]
    ws.append(headers)
    for row in data:
        ws.append([
            row.get("student_id"),
            row.get("student_name"),
            row.get("class_name"),
            row.get("grade"),
            row.get("parent_watch_switch"),
            row.get("openid"),
            row.get("parent_name"),
            row.get("phone"),
            row.get("nickname"),
        ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def export_class_with_students() -> bytes:
    from loongguard.db.parent_db import get_student_parent_bindings
    data = get_student_parent_bindings()
    # 按班级ID排序
    data_sorted = sorted(data, key=lambda x: x.get('class_id') or 0)
    wb = Workbook()
    ws = wb.active
    ws.title = "班级信息"
    headers = ["班级ID", "班级", "学生姓名", "年级", "远程观看权限",
               "家长OpenID", "家长姓名", "家长手机号", "家长昵称"]
    ws.append(headers)
    for row in data_sorted:
        ws.append([
            row.get("class_id"),
            row.get("class_name"),
            row.get("student_name"),
            row.get("grade"),
            row.get("parent_watch_switch"),
            row.get("openid"),
            row.get("parent_name"),
            row.get("phone"),
            row.get("nickname"),
        ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def export_parents_with_students() -> bytes:
    from loongguard.db.parent_db import get_parent_student_bindings
    data = get_parent_student_bindings()
    wb = Workbook()
    ws = wb.active
    ws.title = "家长信息"
    headers = ["家长ID", "家长OpenID", "家长姓名", "家长手机号", "家长昵称",
               "学生ID", "学生姓名", "班级", "年级"]
    ws.append(headers)
    for row in data:
        ws.append([
            row.get("parent_id"),
            row.get("openid"),
            row.get("parent_name"),
            row.get("phone"),
            row.get("nickname"),
            row.get("student_id"),
            row.get("student_name"),
            row.get("class_name"),
            row.get("grade"),
        ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
